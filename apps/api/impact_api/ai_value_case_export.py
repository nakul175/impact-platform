"""Deterministic board/funder PDF and DOCX exports for one complete value case."""

import base64
import hashlib
import io
from copy import deepcopy
from datetime import datetime
from xml.sax.saxutils import escape

from .ai_value_case import DISCLAIMER, calculate_value_case
from .domain import DomainError
from .export_render import normalize_zip

FORMATS = {
    "PDF": ("application/pdf", ".pdf"),
    "DOCX": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".docx",
    ),
}
RENDERER_VERSION = "ai-value-case-export-v1"
MAX_BYTES = 10 * 1024 * 1024


def _invalid(reason="AI_VALUE_CASE_EXPORT_INVALID"):
    raise DomainError("VALIDATION_FAILED", reason=reason)


def _xml_safe(value):
    try:
        value.encode("utf-8")
    except UnicodeError:
        return False
    return all(
        character in "\t\n\r"
        or 0x20 <= ord(character) <= 0xD7FF
        or 0xE000 <= ord(character) <= 0xFFFD
        or 0x10000 <= ord(character) <= 0x10FFFF
        for character in value
    )


def _model(body):
    if not isinstance(body, dict) or set(body) != {"format", "title", "value_case"}:
        _invalid()
    if body["format"] not in FORMATS:
        _invalid("AI_VALUE_CASE_EXPORT_FORMAT_INVALID")
    if not isinstance(body["title"], str) or not body["title"].strip() or len(body["title"]) > 150:
        _invalid()
    result = calculate_value_case(body["value_case"])
    assumptions = deepcopy(result["assumptions"])
    if not assumptions or assumptions != body["value_case"]["assumptions"]:
        _invalid("AI_VALUE_CASE_EXPORT_ASSUMPTIONS_REQUIRED")
    model = {
        "title": body["title"].strip(),
        "format": body["format"],
        "value_case": result,
        "assumptions": assumptions,
        "disclaimer": DISCLAIMER,
    }
    if not all(_xml_safe(text) for text in _all_text(model)):
        _invalid("AI_VALUE_CASE_EXPORT_TEXT_INVALID")
    return model


def _rows(model):
    value = model["value_case"]
    currency = value["currency"]
    rows = [
        ("Opportunity", value["opportunity_id"]),
        ("Planning horizon", str(value["horizon_years"]) + " year(s)"),
        ("Total cost of ownership", currency + " " + value["total_cost_of_ownership"]),
        ("Baseline operating cost", currency + " " + value["baseline_operating_cost"]),
        ("Trial operating cost", currency + " " + value["trial_operating_cost"]),
        ("Gross cost benefit", currency + " " + value["gross_cost_benefit"]),
        ("Net benefit", currency + " " + value["net_benefit"]),
        ("ROI", value["roi_percent"] + "%"),
        ("Cost-benefit ratio", value["cost_benefit_ratio"]),
    ]
    rows.extend(
        (label, currency + " " + value["costs"][key])
        for key, label in (
            ("licences", "Licences"),
            ("setup", "Setup"),
            ("training", "Training"),
            ("staff_time", "Staff time"),
        )
    )
    benefit = value["sector_benefits"]
    rows.extend(
        [
            ("Staff hours saved", benefit["staff_hours_saved"]),
            ("Baseline cost per beneficiary", currency + " " + benefit["baseline_cost_per_beneficiary"]),
            ("Trial cost per beneficiary", currency + " " + benefit["trial_cost_per_beneficiary"]),
            ("Data quality change", benefit["data_quality_change_points"] + " percentage points"),
            ("Reach change", benefit["reach_change"]),
        ]
    )
    return rows


def _all_text(model):
    yield model["title"]
    yield "BOARD OR FUNDER BUSINESS CASE · SELF-REPORTED DRAFT"
    for label, value in _rows(model):
        yield label
        yield value
    for source in ("baseline", "trial"):
        for value in model["value_case"]["source"][source].values():
            yield str(value)
    yield from model["assumptions"]
    yield model["disclaimer"]


def _pdf(model):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    for text in _all_text(model):
        try:
            text.encode("cp1252")
        except UnicodeEncodeError:
            _invalid("AI_VALUE_CASE_EXPORT_GLYPH_UNSUPPORTED")
    styles = getSampleStyleSheet()
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        title=model["title"],
        author="Impact Platform",
        subject="Board or funder business case",
        creator="Impact Platform " + RENDERER_VERSION,
        invariant=1,
    )
    story = [
        Paragraph("BOARD OR FUNDER BUSINESS CASE · SELF-REPORTED DRAFT", styles["Normal"]),
        Paragraph(escape(model["title"]), styles["Title"]),
        Paragraph("Business case", styles["Heading2"]),
    ]
    cell_style = ParagraphStyle("ValueCaseCell", parent=styles["BodyText"], wordWrap="CJK")
    table = Table(
        [[Paragraph("Measure", cell_style), Paragraph("Value", cell_style)]]
        + [
            [Paragraph(escape(label), cell_style), Paragraph(escape(value), cell_style)]
            for label, value in _rows(model)
        ],
        repeatRows=1,
        colWidths=(140, 300),
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5ed")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#ccd6cf")),
            ]
        )
    )
    story.extend([table, Spacer(1, 12), Paragraph("Source inputs", styles["Heading2"])])
    for source in ("baseline", "trial"):
        story.append(Paragraph(source.title(), styles["Heading3"]))
        for key, value in model["value_case"]["source"][source].items():
            story.append(
                Paragraph(escape(key.replace("_", " ").title() + ": " + str(value)), styles["BodyText"])
            )
    story.append(Paragraph("All assumptions", styles["Heading2"]))
    for index, assumption in enumerate(model["assumptions"], 1):
        story.append(Paragraph(escape(str(index) + ". " + assumption), styles["BodyText"]))
    story.extend(
        [
            Spacer(1, 12),
            Paragraph(escape(model["disclaimer"]), styles["Italic"]),
            Paragraph("Rendering " + RENDERER_VERSION, styles["Code"]),
        ]
    )
    document.build(story)
    return output.getvalue()


def _docx(model):
    from docx import Document

    document = Document()
    document.add_paragraph("BOARD OR FUNDER BUSINESS CASE · SELF-REPORTED DRAFT")
    document.add_heading(model["title"], level=0)
    document.add_heading("Business case", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Measure", "Value"
    for label, value in _rows(model):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_heading("Source inputs", level=1)
    for source in ("baseline", "trial"):
        document.add_heading(source.title(), level=2)
        for key, value in model["value_case"]["source"][source].items():
            document.add_paragraph(key.replace("_", " ").title() + ": " + str(value))
    document.add_heading("All assumptions", level=1)
    for assumption in model["assumptions"]:
        document.add_paragraph(assumption, style="List Number")
    document.add_paragraph(model["disclaimer"]).italic = True
    document.add_paragraph("Rendering " + RENDERER_VERSION)
    core = document.core_properties
    core.author = core.last_modified_by = "Impact Platform"
    core.title = model["title"]
    core.subject = "Board or funder business case"
    core.created = core.modified = core.last_printed = datetime(1980, 1, 1)
    core.revision = 1
    output = io.BytesIO()
    document.save(output)
    return normalize_zip(output.getvalue())


def export_value_case(body):
    model = _model(body)
    try:
        raw = _pdf(model) if model["format"] == "PDF" else _docx(model)
    except DomainError:
        raise
    except Exception:
        raise DomainError("SERVICE_UNAVAILABLE", 503, reason="AI_VALUE_CASE_EXPORT_RENDER_FAILED") from None
    if not raw or len(raw) > MAX_BYTES:
        raise DomainError("LIMIT_EXCEEDED", reason="AI_VALUE_CASE_EXPORT_SIZE_LIMIT")
    media_type, extension = FORMATS[model["format"]]
    safe_id = hashlib.sha256(model["value_case"]["opportunity_id"].encode()).hexdigest()[:16]
    return {
        "manifest": {
            "format": model["format"],
            "filename": "impact-value-case-" + safe_id + extension,
            "media_type": media_type,
            "content_sha256": hashlib.sha256(raw).hexdigest(),
            "size_bytes": len(raw),
            "renderer_version": RENDERER_VERSION,
            "assumption_count": len(model["assumptions"]),
        },
        "content_base64": base64.b64encode(raw).decode("ascii"),
    }
