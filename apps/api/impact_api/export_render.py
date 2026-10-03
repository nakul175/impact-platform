"""Deterministic PDF, XLSX and DOCX renderings of a frozen, approved report package (v0.23).

The renderers take one document model built from the immutable revisions a report package binds:
the report revision, its approved template revision, the locked snapshot and the OFFICIAL result
revisions named by its numeric bindings. Every number is the result's stored `displayed_value`
(rounded once when the result was calculated) and its unit, exactly as the HTML and CSV artifacts
of `reporting.py` show it; nothing is recalculated, summed or re-rounded here.

Output is byte-for-byte deterministic for the same model and the same library versions: document
metadata comes from the package (the snapshot's lock time, the report identity), never from the
clock; ZIP containers are rewritten with a fixed entry time and order; the PDF is produced in
reportlab's invariant mode. RENDERER_VERSION names this rendering contract and is recorded with
every artifact, separately from the content (report) version (FR-RPT-003).
"""

import io
import zipfile
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from .report_charts import CHART_NOTE, chart_description, chart_drawing, chart_models, chart_rows
from .reporting import BINDING_TOKEN

RENDERER_VERSION = "exports-1"
CHART_COLUMNS = ["bar", "official_value", "unit", "target", "value_state"]
MEDIA_TYPES = {
    "PDF": "application/pdf",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
EXTENSIONS = {"PDF": ".pdf", "XLSX": ".xlsx", "DOCX": ".docx"}
BANNER = "CONTROLLED · APPROVED SNAPSHOT PACKAGE"
TABLE_COLUMNS = [
    "section_code",
    "section_heading",
    "binding_code",
    "displayed_value",
    "unit",
    "result_revision",
    "report_revision",
    "snapshot_id",
]
# The largest artifact stored (the migration's CHECK enforces the same bound).
MAX_ARTIFACT_BYTES = 10 * 1024 * 1024
ZIP_TIME = (1980, 1, 1, 0, 0, 0)


class RenderError(Exception):
    """A rendering failure: only the error class is recorded, never content."""

    def __init__(self, error_class, permanent=True):
        super().__init__(error_class)
        self.error_class = error_class
        self.permanent = permanent


def display(result):
    """The value and unit exactly as the HTML artifact prints them."""
    return str(result.get("displayed_value", "—")), str(result.get("unit", ""))


def build_model(
    report_id, report_revision, report, template, snapshot_id, snapshot, digest_hex, values, targets=None
):
    """`values` maps (section_code, binding_code) to the bound OFFICIAL result payload; `targets`
    maps a pinned target revision to its payload (charts, v0.27)."""
    sections, rows = [], []
    for section in report["sections"]:
        paragraphs = []
        for paragraph in section["narrative"].splitlines() or [""]:
            runs, start = [], 0
            for match in BINDING_TOKEN.finditer(paragraph):
                if match.start() > start:
                    runs.append(("text", paragraph[start : match.start()]))
                value, unit = display(values[(section["section_code"], match.group(1))])
                runs.append(("bound", match.group(1), value, unit))
                start = match.end()
            if start < len(paragraph) or not runs:
                runs.append(("text", paragraph[start:]))
            paragraphs.append(runs)
        table = []
        for numeric in section["bindings"]:
            value, unit = display(values[(section["section_code"], numeric["binding_code"])])
            table.append((numeric["binding_code"], value, unit))
            rows.append(
                [
                    section["section_code"],
                    section["heading"],
                    numeric["binding_code"],
                    value,
                    unit,
                    numeric["result_revision"],
                    str(report_revision),
                    str(snapshot_id),
                ]
            )
        sections.append(
            {
                "heading": section["heading"],
                "paragraphs": paragraphs,
                "table": table,
                "charts": chart_models(section, values, targets),
            }
        )
    locked_at = snapshot["locked_at"]
    return {
        "title": template.get("title", "Impact report"),
        "language": report["language"],
        "sections": sections,
        "rows": rows,
        "footer": {
            "report_id": str(report_id),
            "report_revision": str(report_revision),
            "snapshot_id": str(snapshot_id),
            "locked_at": locked_at,
            "reconciliation_sha256": digest_hex,
        },
    }


def model_instant(model):
    """The snapshot's lock time as a naive UTC datetime: the only time any artifact carries."""
    value = datetime.fromisoformat(model["footer"]["locked_at"].replace("Z", "+00:00"))
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).replace(tzinfo=None, microsecond=0)


def footer_lines(model):
    f = model["footer"]
    return [
        "Report " + f["report_id"] + " · revision " + f["report_revision"],
        "Snapshot " + f["snapshot_id"] + " · locked " + f["locked_at"],
        "Reconciliation SHA-256 " + f["reconciliation_sha256"],
        "Rendering " + RENDERER_VERSION,
    ]


def paragraph_text(runs):
    return "".join(run[1] if run[0] == "text" else run[2] + " " + run[3] for run in runs)


def all_text(model):
    yield model["title"]
    yield BANNER
    for section in model["sections"]:
        yield section["heading"]
        for runs in section["paragraphs"]:
            yield paragraph_text(runs)
        for row in section["table"]:
            yield " ".join(row)
        for chart in section.get("charts", []):
            yield chart["title"]
            yield chart_description(chart)
    yield from footer_lines(model)


def normalize_zip(data):
    """Rewrite a ZIP container with a fixed entry time, permissions and order, so the bytes depend
    only on the entries' content."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            info.create_system = 0
            target.writestr(info, source.read(name))
    return output.getvalue()


def render_pdf(model):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    # The standard Helvetica faces cover Windows-1252 only; a glyph they cannot draw would be
    # silently lost, so such a package is refused rather than rendered incompletely.
    for text in all_text(model):
        try:
            text.encode("cp1252")
        except UnicodeEncodeError:
            raise RenderError("RENDER_GLYPH_UNSUPPORTED") from None
    styles = getSampleStyleSheet()
    buffer = io.BytesIO()
    f = model["footer"]
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        title=model["title"],
        author="Impact Platform",
        subject="Report " + f["report_id"] + " revision " + f["report_revision"],
        creator="Impact Platform " + RENDERER_VERSION,
        keywords="snapshot " + f["snapshot_id"] + "; reconciliation " + f["reconciliation_sha256"],
        lang=model["language"],
        invariant=1,
    )
    story = [
        Paragraph(escape(BANNER), styles["Normal"]),
        Paragraph(escape(model["title"]), styles["Title"]),
    ]
    for section in model["sections"]:
        story.append(Paragraph(escape(section["heading"]), styles["Heading2"]))
        for runs in section["paragraphs"]:
            markup = "".join(
                escape(run[1]) if run[0] == "text" else "<b>" + escape(run[2] + " " + run[3]) + "</b>"
                for run in runs
            )
            story.append(Paragraph(markup, styles["BodyText"]))
        if section["table"]:
            data = [["Metric", "Value and unit"]] + [
                [code, value + " " + unit] for code, value, unit in section["table"]
            ]
            table = Table(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTNAME", (1, 1), (1, -1), "Helvetica-Bold"),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5ed")),
                        ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#ccd6cf")),
                    ]
                )
            )
            story += [Spacer(1, 6), Paragraph("Bound official results", styles["Italic"]), table]
        for chart in section.get("charts", []):
            data = [["Bar", "Official value", "Unit", "Target", "Value state"]] + chart_rows(chart)
            table = Table(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5ed")),
                        ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#ccd6cf")),
                    ]
                )
            )
            story += [
                Spacer(1, 12),
                chart_drawing(chart),
                Paragraph(escape(chart["title"]) + ". " + escape(CHART_NOTE), styles["Italic"]),
                table,
            ]
    story.append(Spacer(1, 18))
    story += [Paragraph(escape(line), styles["Code"]) for line in footer_lines(model)]
    document.build(story)
    return buffer.getvalue()


def text_cell(sheet, row, column, value):
    """Every cell is stored as text: a value beginning with '=' is never a formula, and a displayed
    number is never converted, re-rounded or summed by the spreadsheet."""
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    cell.data_type = "s"
    return cell


def render_xlsx(model):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.writer.excel import ExcelWriter

    workbook = Workbook()
    values = workbook.active
    values.title = "Bound values"
    for column, name in enumerate(TABLE_COLUMNS, 1):
        text_cell(values, 1, column, name).font = Font(bold=True)
    for index, row in enumerate(model["rows"], 2):
        for column, value in enumerate(row, 1):
            text_cell(values, index, column, value)
    values.freeze_panes = "A2"
    charts = [chart for section in model["sections"] for chart in section.get("charts", [])]
    if charts:
        sheet = workbook.create_sheet("Charts")
        text_cell(sheet, 1, 1, CHART_NOTE)
        line = 3
        for chart in charts:
            text_cell(sheet, line, 1, chart["title"]).font = Font(bold=True)
            line += 1
            for column, name in enumerate(CHART_COLUMNS, 1):
                text_cell(sheet, line, column, name).font = Font(bold=True)
            line += 1
            for row in chart_rows(chart):
                for column, value in enumerate(row, 1):
                    text_cell(sheet, line, column, value)
                line += 1
            line += 1
    package = workbook.create_sheet("Package")
    lines = [("title", model["title"]), ("classification", BANNER), ("language", model["language"])]
    for section in model["sections"]:
        lines.append(("section", section["heading"]))
        lines += [("narrative", paragraph_text(runs)) for runs in section["paragraphs"]]
    lines += [(key, value) for key, value in model["footer"].items()]
    lines.append(("renderer_version", RENDERER_VERSION))
    for index, (key, value) in enumerate(lines, 1):
        text_cell(package, index, 1, key).font = Font(bold=True)
        text_cell(package, index, 2, value)
    instant = model_instant(model)
    properties = workbook.properties
    properties.creator = properties.lastModifiedBy = "Impact Platform"
    properties.title = model["title"]
    properties.subject = (
        "Report " + model["footer"]["report_id"] + " revision " + model["footer"]["report_revision"]
    )
    properties.created = properties.modified = instant
    properties.language = model["language"]
    buffer = io.BytesIO()
    # ExcelWriter directly: openpyxl's save_workbook would stamp the current time as `modified`.
    archive = zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, allowZip64=True)
    ExcelWriter(workbook, archive).save()
    return normalize_zip(buffer.getvalue())


def render_docx(model):
    from docx import Document

    document = Document()
    document.add_paragraph(BANNER)
    document.add_heading(model["title"], level=0)
    for section in model["sections"]:
        document.add_heading(section["heading"], level=1)
        for runs in section["paragraphs"]:
            paragraph = document.add_paragraph()
            for run in runs:
                if run[0] == "text":
                    paragraph.add_run(run[1])
                else:
                    paragraph.add_run(run[2] + " " + run[3]).bold = True
        if section["table"]:
            document.add_paragraph("Bound official results").italic = True
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            header = table.rows[0].cells
            header[0].text, header[1].text = "Metric", "Value and unit"
            for code, value, unit in section["table"]:
                cells = table.add_row().cells
                cells[0].text = code
                cells[1].paragraphs[0].add_run(value + " " + unit).bold = True
        for chart in section.get("charts", []):
            document.add_paragraph(chart["title"] + ". " + CHART_NOTE).italic = True
            table = document.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            for cell, name in zip(
                table.rows[0].cells, ["Bar", "Official value", "Unit", "Target", "Value state"]
            ):
                cell.text = name
            for row in chart_rows(chart):
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
    for line in footer_lines(model):
        document.add_paragraph(line)
    instant = model_instant(model)
    core = document.core_properties
    core.author = core.last_modified_by = "Impact Platform"
    core.title = model["title"]
    core.subject = (
        "Report " + model["footer"]["report_id"] + " revision " + model["footer"]["report_revision"]
    )
    core.language = model["language"]
    core.created = core.modified = core.last_printed = instant
    core.revision = 1
    core.comments = "Rendering " + RENDERER_VERSION
    buffer = io.BytesIO()
    document.save(buffer)
    return normalize_zip(buffer.getvalue())


RENDERERS = {"PDF": render_pdf, "XLSX": render_xlsx, "DOCX": render_docx}


def render(format, model):
    if format not in RENDERERS:
        raise RenderError("FORMAT_UNSUPPORTED")
    try:
        body = RENDERERS[format](model)
    except RenderError:
        raise
    except Exception:  # a library defect is a permanent failure of this rendering, never a crash
        raise RenderError("RENDER_FAILED") from None
    if not body or len(body) > MAX_ARTIFACT_BYTES:
        raise RenderError("EXPORT_ARTIFACT_LIMIT")
    return body
