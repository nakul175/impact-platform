"""Approved framework exports; no current targets, observations or recalculated values.

Workbook header styling is adapted from Mercy Corps TolaData xls_export_utils.py,
commit 7ca89ab1e5f55cbe4577d16d7281c6cf0936fc3d (Apache-2.0).
See third_party/mercycorps-toladata for license and attribution. Hierarchy and export
semantics are implemented here for Impact Platform's immutable framework revisions.
"""

import csv
import hashlib
import io
import zipfile
from datetime import datetime

from .domain import DomainError, unavailable

MEDIA = {
    "CSV": "text/csv; charset=utf-8",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
VERSION = "logframe-1"


def safe_text(value):
    """Spreadsheet formula guard, including formulas preceded by whitespace/control characters."""
    text = "" if value is None else str(value)
    leading = "".join(chr(i) for i in range(33)) + "\ufeff"
    return "'" + text if text.lstrip(leading).startswith(("=", "+", "-", "@")) else text


def tables(framework_id, revision_id, data):
    """Stable preorder preserves authored sibling order, even when children precede parents."""
    nodes = data.get("nodes", [])
    by_id = {n["node_id"]: n for n in nodes}
    if len(by_id) != len(nodes) or len(nodes) > 500:
        raise DomainError("INVALID_STATE", 409, reason="INVALID_FRAMEWORK_STRUCTURE")
    children = {}
    for node in nodes:
        parent = node.get("parent_node_id")
        if parent is not None and parent not in by_id:
            raise DomainError("INVALID_STATE", 409, reason="INVALID_FRAMEWORK_STRUCTURE")
        children.setdefault(parent, []).append(node)
    ordered, seen = [], set()
    stack = [(n, 0) for n in reversed(children.get(None, []))]
    while stack:
        node, depth = stack.pop()
        if node["node_id"] in seen:
            raise DomainError("INVALID_STATE", 409, reason="INVALID_FRAMEWORK_STRUCTURE")
        seen.add(node["node_id"])
        ordered.append(
            [
                depth,
                node["node_id"],
                node.get("parent_node_id"),
                node["node_type"],
                node["title"],
                node["definition"],
                node.get("owner_id"),
                " | ".join(node.get("indicator_ids", [])),
            ]
        )
        stack.extend((n, depth + 1) for n in reversed(children.get(node["node_id"], [])))
    if len(seen) != len(nodes):
        raise DomainError("INVALID_STATE", 409, reason="INVALID_FRAMEWORK_STRUCTURE")
    return {
        "Context": [
            ["field", "value"],
            ["framework_id", framework_id],
            ["framework_revision", revision_id],
            ["programme_id", data["programme_id"]],
            ["version_label", data.get("version_label")],
            ["effective_from", data.get("effective_from")],
            ["renderer_version", VERSION],
            ["scope", "Approved framework structure only; indicator identifiers, no measurement values."],
        ],
        "Logframe": [
            [
                "depth",
                "node_id",
                "parent_node_id",
                "level",
                "title",
                "definition",
                "owner_id",
                "indicator_ids",
            ],
            *ordered,
        ],
        "Relationships": [
            [
                "relationship_id",
                "from_node_id",
                "to_node_id",
                "type",
                "rationale",
                "evidence_strength",
                "assumption_ids",
                "external_context",
            ],
            *[
                [
                    r["relationship_id"],
                    r["from_node_id"],
                    r["to_node_id"],
                    r["relationship_type"],
                    r["rationale"],
                    r["evidence_strength"],
                    " | ".join(r.get("assumption_ids", [])),
                    r.get("external_context"),
                ]
                for r in data.get("relationships", [])
            ],
        ],
        "Assumptions": [
            [
                "assumption_id",
                "kind",
                "node_ids",
                "statement",
                "expected_condition",
                "evidence",
                "owner_id",
                "review_date",
                "status",
            ],
            *[
                [
                    a["assumption_id"],
                    a["kind"],
                    " | ".join(a.get("node_ids", [])),
                    a["statement"],
                    a.get("expected_condition"),
                    a.get("evidence"),
                    a.get("owner_id"),
                    a["review_date"],
                    a["status"],
                ]
                for a in data.get("assumptions", [])
            ],
        ],
        "Exceptions": [
            ["object_id", "rule", "reason", "review_date", "recorded_by", "recorded_at"],
            *[
                [
                    e["object_id"],
                    e["rule"],
                    e["reason"],
                    e["review_date"],
                    e.get("recorded_by"),
                    e.get("recorded_at"),
                ]
                for e in data.get("exceptions", [])
            ],
        ],
    }


def render(model, format):
    if format == "CSV":
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n")
        # A tagged, rectangular representation preserves every workbook section in one CSV.
        writer.writerow(["section", *[f"column_{i}" for i in range(1, 10)]])
        for section, rows in model.items():
            for row in rows:
                writer.writerow([section, *[safe_text(v) for v in row], *[""] * (9 - len(row))])
        return buffer.getvalue().encode("utf-8-sig")
    if format != "XLSX":
        raise DomainError("VALIDATION_FAILED", reason="LOGFRAME_FORMAT_INVALID")
    from openpyxl import Workbook
    from openpyxl.writer.excel import ExcelWriter
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Impact Platform"
    workbook.properties.created = workbook.properties.modified = datetime(2000, 1, 1)
    for section, rows in model.items():
        sheet = workbook.create_sheet(section)
        for row in rows:
            values = [safe_text(v) for v in row]
            if any(ILLEGAL_CHARACTERS_RE.search(value) for value in values):
                raise DomainError("VALIDATION_FAILED", reason="LOGFRAME_TEXT_NOT_SUPPORTED")
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"  # identifiers and decimal-looking text never become formulas/numbers
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        # Adapted from apply_label_styling; platform palette, no Mercy Corps branding.
        for cell in sheet[1]:
            cell.font = Font(bold=True, size=10)
            cell.fill = PatternFill(fill_type="solid", start_color="E8EEF4", end_color="E8EEF4")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            edge = Side(style="thin", color="243746")
            cell.border = Border(left=edge, top=edge, right=edge, bottom=edge)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28
    raw = io.BytesIO()
    # save_workbook stamps the wall clock; ExcelWriter preserves our fixed metadata.
    with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(workbook, archive).save()
    result = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            info.create_system = 0
            target.writestr(info, source.read(name))
    return result.getvalue()


class LogframeExports:
    def __init__(self, service):
        self.service = service

    def download(self, identity, tenant, framework_id, revision_id, format, correlation):
        from .service import revision
        from .store import audit, authorize, context, load

        if format not in MEDIA:
            raise DomainError("VALIDATION_FAILED", reason="LOGFRAME_FORMAT_INVALID")
        with self.service.db.transaction(tenant) as c:
            c.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))", (tenant,))
            ctx = context(c, identity, tenant, write=True)
            authorize(c, ctx, "export_framework_" + format.lower(), framework_id, hidden=True)
            pinned = revision(c, ctx, revision_id, "Framework", "frameworks.read")
            if str(pinned["object_id"]) != framework_id:
                unavailable()
            approved = c.execute(
                "SELECT approved_at FROM impact.framework_baseline WHERE tenant_id=%s "
                "AND framework_id=%s AND framework_revision=%s",
                (tenant, framework_id, revision_id),
            ).fetchone()
            if not approved:
                raise DomainError("INVALID_STATE", 409, reason="APPROVED_FRAMEWORK_REQUIRED")
            data = pinned["payload"]
            load(c, ctx, data["programme_id"], "Programme", "programmes.read")
            # Refuse the whole export if any placed indicator is currently hidden/restricted.
            for node in data.get("nodes", []):
                for indicator in node.get("indicator_ids", []):
                    load(c, ctx, indicator, "IndicatorInstance", "indicator-instances.read")
            body = render(tables(framework_id, revision_id, data), format)
            if len(body) > 10 * 1024 * 1024:
                raise DomainError("LIMIT_EXCEEDED", 422, reason="LOGFRAME_EXPORT_LIMIT")
            at = c.execute("SELECT statement_timestamp() AS at").fetchone()["at"]
            audit(
                c,
                ctx,
                "export_framework_" + format.lower(),
                {
                    "object_id": framework_id,
                    "revision_id": revision_id,
                    "business_state": "Approved",
                    "saved_at": at.isoformat(),
                },
                correlation,
            )
            return body, hashlib.sha256(body).hexdigest()
