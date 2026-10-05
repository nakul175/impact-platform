"""Mercy Corps inspired hierarchy cases with our immutable export and authority boundary."""
# ruff: noqa: F811

import csv
import io
import zipfile
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from impact_api.domain import DomainError
from impact_api.logframe import render, safe_text, tables
from test_measurement import approve, get, setup, submit  # noqa: F401
from test_planning import framework


def node(id, parent=None, **extra):
    return {
        "node_id": id,
        "parent_node_id": parent,
        "node_type": "OUTPUT",
        "title": "Éducation शिक्षा",
        "definition": "Intended change",
        "indicator_ids": [],
        **extra,
    }


def payload(nodes):
    return {"programme_id": "programme", "version_label": "Baseline", "nodes": nodes}


def test_hierarchy_orders_parents_first_and_preserves_sibling_order():
    model = tables("f", "r", payload([node("b", "a"), node("a"), node("c", "a")]))
    assert [(r[0], r[1]) for r in model["Logframe"][1:]] == [(0, "a"), (1, "b"), (1, "c")]


@pytest.mark.parametrize(
    "nodes", [[node("a", "b")], [node("a", "b"), node("b", "a")], [node("a"), node("a")]]
)
def test_invalid_structure_fails_closed(nodes):
    with pytest.raises(DomainError) as exc:
        tables("f", "r", payload(nodes))
    assert exc.value.reason == "INVALID_FRAMEWORK_STRUCTURE"


def test_empty_and_deep_frameworks():
    assert len(tables("f", "r", payload([]))["Logframe"]) == 1
    nodes = [node(str(i), str(i - 1) if i else None) for i in range(500)]
    assert tables("f", "r", payload(nodes))["Logframe"][-1][0] == 499


@pytest.mark.parametrize("text", ["=1+1", " +SUM(A1)", "\t@SUM(A1)", "\n-1+1", "\ufeff=1+1", "\x00=1+1"])
def test_formula_guard(text):
    assert safe_text(text) == "'" + text


def test_workbook_preserves_unicode_context_and_formula_safety():
    model = tables("f", "r", payload([node("a", definition='=HYPERLINK("x")')]))
    data = render(model, "XLSX")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Context", "Logframe", "Relationships", "Assumptions", "Exceptions"]
    assert wb["Logframe"]["E2"].value == "Éducation शिक्षा"
    assert wb["Logframe"]["F2"].value.startswith("'=HYPERLINK")
    assert wb["Logframe"]["F2"].data_type == "s"
    assert wb["Context"]["B3"].value == "r"
    assert wb["Logframe"].freeze_panes == "A2"


def test_workbook_refuses_unrepresentable_control_characters():
    with pytest.raises(DomainError) as exc:
        render(tables("f", "r", payload([node("a", title="\x00=1+1")])), "XLSX")
    assert exc.value.reason == "LOGFRAME_TEXT_NOT_SUPPORTED"


def test_deterministic_bytes_and_no_clock_metadata():
    model = tables("f", "r", payload([node("a")]))
    for format in ["CSV", "XLSX"]:
        assert render(model, format) == render(model, format)
    with zipfile.ZipFile(io.BytesIO(render(model, "XLSX"))) as archive:
        assert all(i.date_time == (1980, 1, 1, 0, 0, 0) for i in archive.infolist())
        assert b"2000-01-01T00:00:00Z" in archive.read("docProps/core.xml")


def test_csv_is_rectangular_and_preserves_all_sections():
    data = payload([node("a", title="=1+1")])
    data["exceptions"] = [
        {"object_id": "a", "rule": "UNMEASURED_RESULT", "reason": "Endline", "review_date": "2027-01-01"}
    ]
    rows = list(csv.reader(io.StringIO(render(tables("f", "r", data), "CSV").decode("utf-8-sig"))))
    assert {len(r) for r in rows} == {10}
    assert {r[0] for r in rows[1:]} == {"Context", "Logframe", "Relationships", "Assumptions", "Exceptions"}
    assert any("'=1+1" in r for r in rows)


def download(live, framework, revision=None, actor="author", format="csv"):
    return live.request(
        live.path("frameworks", framework["object_id"] + "/logframe." + format),
        actor=actor,
        params={"revision_id": revision or framework["revision_id"]},
    )


def test_approved_export_is_pinned_audited_and_rejects_unapproved(live, setup):
    programme, indicator, _, _ = setup(False)
    draft = framework(live, programme, indicator)
    assert download(live, draft).status_code == 409
    approve(live, submit(live, "frameworks", draft))
    approved = get(live, "frameworks", draft["object_id"])
    first = download(live, approved)
    assert first.status_code == 200
    assert approved["revision_id"] in first.text
    assert download(live, approved).content == first.content
    assert download(live, approved, format="xlsx").status_code == 200
    # A newly-created, changed draft cannot alter the bytes of the pinned approved baseline.
    changed = deepcopy(approved["data"])
    changed["version_label"] = "Later draft"
    from test_measurement import create

    for exc in changed.get("exceptions", []):
        exc.pop("recorded_by", None)
        exc.pop("recorded_at", None)
    changed["supersedes_revision"] = approved["revision_id"]
    changed["effective_from"] = "2026-02-01T00:00:00Z"
    child = create(live, "frameworks", changed)
    assert download(live, approved, revision=child["revision_id"]).status_code == 404
    approve(live, submit(live, "frameworks", child))
    assert get(live, "frameworks", approved["object_id"])["lifecycle_state"] == "Superseded"
    assert download(live, approved).content == first.content
    with live.db() as c:
        rows = c.execute(
            "SELECT payload FROM impact.object_revision WHERE object_type='AuditEvent' "
            "AND payload->>'action_type'='export_framework_csv'"
        ).fetchall()
    assert any(r["payload"]["object_reference"] == approved["object_id"] for r in rows)


def test_export_requires_separate_authority_and_current_source_visibility(live, setup):
    programme, indicator, _, _ = setup(False)
    draft = framework(live, programme, indicator)
    approve(live, submit(live, "frameworks", draft))
    approved = get(live, "frameworks", draft["object_id"])
    # Remove only the fixture reviewer's export authority; read stays available.
    assert get(live, "frameworks", approved["object_id"], actor="reviewer")
    principal = live.fixture["actors"]["reviewer"]["principal_id"]
    with live.db() as c:
        grants = c.execute(
            "UPDATE impact.grant_current SET purpose='QUALIFICATION' "
            "WHERE subject_id=%s AND capability='framework.export' AND purpose IS NULL RETURNING object_id",
            (principal,),
        ).fetchall()
    try:
        assert download(live, approved, actor="reviewer").status_code == 404
        assert get(live, "frameworks", approved["object_id"], actor="reviewer")
    finally:
        with live.db() as c:
            c.execute(
                "UPDATE impact.grant_current SET purpose=NULL WHERE object_id=ANY(%s::uuid[])",
                ([str(g["object_id"]) for g in grants],),
            )
    for actor in ["other_tenant", "revoked"]:
        assert download(live, approved, actor=actor).status_code in {401, 404}
    with live.db() as c:
        old = c.execute(
            "SELECT classification FROM impact.object_registry WHERE object_id=%s", (indicator["object_id"],)
        ).fetchone()["classification"]
        c.execute(
            "UPDATE impact.object_registry SET classification='RESTRICTED' WHERE object_id=%s",
            (indicator["object_id"],),
        )
    try:
        assert download(live, approved).status_code == 404
    finally:
        with live.db() as c:
            c.execute(
                "UPDATE impact.object_registry SET classification=%s WHERE object_id=%s",
                (old, indicator["object_id"]),
            )
