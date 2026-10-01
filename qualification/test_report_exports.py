"""Report exports (v0.23): PDF, XLSX and DOCX renderings of an approved, frozen report package,
requested as a business command, executed by the worker under generation-fenced leases, stored
immutably and delivered through controlled publication. The worker runs in-process one iteration
or one step at a time (as in test_worker.py); natively it uses the provisioned worker login."""
# ruff: noqa: F811

import base64
import hashlib
import io
import re
import threading
import uuid
import zlib
from datetime import timedelta

import psycopg
import pytest
from psycopg.rows import dict_row

from impact_api.export_render import (
    MEDIA_TYPES,
    RENDERER_VERSION,
    RenderError,
    build_model,
    render,
)
from impact_api.worker import empty_summary
from test_live_application import cmd, expect
from test_measurement import get, create, action, submit, approve  # noqa: F401
from test_publication import approved_report, publish, request_disclosure
from test_reporting import package_data
from test_worker import NATIVE, Clock, make_worker, worker_dsn

FORMATS = ["PDF", "XLSX", "DOCX"]


# -- helpers ---------------------------------------------------------------------------------------


def request_export(live, report, format, actor="author", status=200, body=None):
    return expect(
        live.request(
            live.path("reports", report["object_id"]) + "/actions/export",
            actor=actor,
            method="POST",
            body=body or cmd({"format": format}, report["revision_id"]),
        ),
        status,
    )


def exports(live, report, actor="author"):
    return expect(live.request(live.path("reports", report["object_id"]) + "/exports", actor=actor), 200)[
        "items"
    ]


def export_row(live, report, job_id):
    return next(item for item in exports(live, report) if item["job_id"] == job_id)


def job(live, job_id):
    with live.db() as c:
        return c.execute(
            "SELECT j.*,e.attempts,e.lease_owner,e.next_attempt_at,e.last_error_class,e.completed_at "
            "FROM impact.job j JOIN impact.report_export e USING(tenant_id,job_id) WHERE j.job_id=%s",
            (job_id,),
        ).fetchone()


def artifacts(live, job_id):
    with live.db() as c:
        return c.execute("SELECT * FROM impact.report_export_artifact WHERE job_id=%s", (job_id,)).fetchall()


def drain(worker, tenant, summary=None):
    """Claim and process every due export of the tenant with this worker."""
    summary = summary or empty_summary()
    for _ in range(20):
        rows = worker.claim_exports(tenant, summary)
        if not rows:
            break
        for row in rows:
            worker.process_export(tenant, row, summary)
    return summary


def pdf_text(body):
    """The text drawing operators of a reportlab PDF (ASCII85 + Flate content streams)."""
    text = []
    for match in re.finditer(rb"stream\r?\n(.*?)endstream", body, re.S):
        raw = match.group(1).strip()
        if raw.endswith(b"~>"):
            raw = base64.a85decode(raw[:-2])
        try:
            text.append(zlib.decompress(raw).decode("latin-1"))
        except zlib.error:
            continue
    return "\n".join(text)


def xlsx_rows(body):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body))
    return {
        sheet.title: [[(cell.value, cell.data_type) for cell in row] for row in sheet.iter_rows()]
        for sheet in workbook
    }


def docx_text(body):
    from docx import Document

    document = Document(io.BytesIO(body))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return lines


def bound_result(live):
    """The OFFICIAL result revision the package binds, read from the immutable revision."""
    revision = package_data(live)["sections"][0]["bindings"][0]["result_revision"]
    with live.db() as c:
        return c.execute(
            "SELECT payload FROM impact.object_revision WHERE revision_id=%s", (revision,)
        ).fetchone()["payload"]


def settle(live):
    """Finish every export earlier tests left due, so a test's claims see its own jobs."""
    drain(make_worker(live), live.fixture["tenant_a"])


def produced(live, formats=FORMATS):
    """An approved report with one succeeded export per format; returns (report, {format: job})."""
    report = approved_report(live)
    jobs = {format: request_export(live, report, format)["job_id"] for format in formats}
    drain(make_worker(live), live.fixture["tenant_a"])
    for format, job_id in jobs.items():
        assert export_row(live, report, job_id)["state"] == "Succeeded", format
    return report, jobs


# -- pure rendering --------------------------------------------------------------------------------


def model(**changes):
    report = {
        "language": "en",
        "sections": [
            {
                "section_code": "results",
                "heading": '=HYPERLINK("http://x")',
                "narrative": "Coverage reached {{water}} this quarter.\nSecond paragraph.",
                "bindings": [{"binding_code": "water", "result_revision": str(uuid.uuid4())}],
            }
        ],
    }
    report.update(changes)
    return build_model(
        "r-1",
        "rev-1",
        report,
        {"title": "Quarterly report"},
        "s-1",
        {"locked_at": "2026-09-30T10:00:00+00:00"},
        "ab" * 32,
        {("results", "water"): {"displayed_value": "46.36", "unit": "PERCENT"}},
    )


def test_renderers_are_byte_deterministic_and_carry_only_displayed_values():
    document = model()
    for format in FORMATS:
        first, second = render(format, document), render(format, document)
        assert hashlib.sha256(first).digest() == hashlib.sha256(second).digest(), format
    pdf = pdf_text(render("PDF", document))
    assert "(46.36 PERCENT) Tj" in pdf and "(Coverage reached ) Tj" in pdf
    sheets = xlsx_rows(render("XLSX", document))
    header, row = sheets["Bound values"]
    assert [value for value, _ in header][3:5] == ["displayed_value", "unit"]
    # Text cells only: the displayed value is never converted, and '=' never becomes a formula.
    assert row[3] == ("46.36", "s") and row[1] == ('=HYPERLINK("http://x")', "s")
    assert all(kind == "s" for lines in sheets.values() for line in lines for _, kind in line)
    lines = docx_text(render("DOCX", document))
    assert "Coverage reached 46.36 PERCENT this quarter." in lines
    assert "water | 46.36 PERCENT" in lines


def test_pdf_refuses_glyphs_it_cannot_draw_instead_of_dropping_them():
    document = model(language="hi")
    document["sections"][0]["heading"] = "जल आपूर्ति"
    with pytest.raises(RenderError) as refused:
        render("PDF", document)
    assert refused.value.error_class == "RENDER_GLYPH_UNSUPPORTED" and refused.value.permanent
    # XLSX and DOCX carry Unicode text natively.
    assert render("XLSX", document) and render("DOCX", document)


# -- request ---------------------------------------------------------------------------------------


def test_export_request_is_one_queued_job_with_receipt_audit_and_exact_retry(live):
    report = approved_report(live)
    body = cmd({"format": "PDF"}, report["revision_id"])
    receipt = request_export(live, report, "PDF", body=body)
    assert receipt["business_state"] == "Queued" and receipt["object_id"] == report["object_id"]
    assert receipt["revision_id"] == report["revision_id"]
    assert request_export(live, report, "PDF", body=body) == receipt
    changed = {**body, "data": {"format": "DOCX"}}
    assert request_export(live, report, "DOCX", body=changed, status=409)["code"] == "CONFLICT_OPERATION"
    duplicate = request_export(live, report, "PDF", status=409)
    assert duplicate["reason_code"] == "EXPORT_ALREADY_REQUESTED"
    stale = request_export(live, report, "XLSX", body=cmd({"format": "XLSX"}, str(uuid.uuid4())), status=409)
    assert stale["code"] == "CONFLICT_VERSION"

    row = job(live, receipt["job_id"])
    assert (row["job_class"], row["state"], row["lease_generation"], row["attempts"]) == (
        "REPORT_EXPORT",
        "Queued",
        0,
        0,
    )
    assert row["input_manifest"]["report_revision"] == report["revision_id"]
    assert row["input_manifest"]["format"] == "PDF"
    assert row["input_manifest"]["renderer_version"] == RENDERER_VERSION
    assert str(row["requester_id"]) == live.fixture["actors"]["author"]["principal_id"]
    with live.db() as c:
        binding = c.execute(
            "SELECT reconciliation_digest FROM impact.report_package_binding WHERE report_revision=%s",
            (report["revision_id"],),
        ).fetchone()
        assert row["input_manifest"]["reconciliation_sha256"] == bytes(binding["reconciliation_digest"]).hex()
        assert (
            c.execute(
                "SELECT count(*) AS n FROM impact.operation_receipt WHERE operation_id=%s",
                (body["operation_id"],),
            ).fetchone()["n"]
            == 1
        )
        assert (
            c.execute(
                "SELECT count(*) AS n FROM impact.audit_event_current WHERE action_type='action_reports_export' "
                "AND correlation_id=%s AND object_reference=%s",
                (receipt["correlation_id"], report["object_id"]),
            ).fetchone()["n"]
            == 1
        )
    [listed] = [item for item in exports(live, report) if item["job_id"] == receipt["job_id"]]
    assert listed["state"] == "Queued" and listed["content_sha256"] is None


def test_export_request_needs_an_approved_package_the_capability_and_the_tenant(live):
    draft = create(live, "reports", package_data(live))
    assert request_export(live, draft, "PDF", status=409)["reason_code"] == "APPROVED_REPORT_REQUIRED"
    report = approved_report(live)
    # EXTERNAL holds no report.export; another tenant and a revoked member cannot see the report.
    assert request_export(live, report, "PDF", actor="partner", status=404)["code"] == "RESOURCE_UNAVAILABLE"
    assert request_export(live, report, "PDF", actor="other_tenant", status=404)
    assert live.request(
        live.path("reports", report["object_id"]) + "/actions/export",
        actor="revoked",
        method="POST",
        body=cmd({"format": "PDF"}, report["revision_id"]),
    ).status_code in {401, 404}
    expect(live.request(live.path("reports", report["object_id"]) + "/exports", actor="partner"), 404)
    expect(live.request(live.path("reports", report["object_id"]) + "/exports", actor="other_tenant"), 404)
    unknown = request_export(live, report, "PPTX", status=422)
    assert unknown["code"] == "VALIDATION_FAILED"


# -- execution -------------------------------------------------------------------------------------


def test_worker_renders_immutable_artifacts_whose_numbers_match_the_snapshot_and_html(live):
    report, jobs = produced(live)
    result = bound_result(live)
    displayed = result["displayed_value"] + " " + result["unit"]
    assert result["mode"] == "OFFICIAL" and displayed == "46.36 PERCENT"
    html = live.request(live.path("reports", report["object_id"]) + "/export").text
    csv_text = live.request(live.path("reports", report["object_id"]) + "/export.csv").text
    assert 'data-binding="water">' + displayed in html
    with live.db() as c:
        digest = bytes(
            c.execute(
                "SELECT reconciliation_digest FROM impact.report_package_binding WHERE report_revision=%s",
                (report["revision_id"],),
            ).fetchone()["reconciliation_digest"]
        ).hex()
    assert digest in html

    bodies = {}
    for format, job_id in jobs.items():
        listed = export_row(live, report, job_id)
        response = live.request(
            live.path("reports", report["object_id"]) + "/exports/" + job_id + "/download"
        )
        assert response.status_code == 200
        body = response.content
        bodies[format] = body
        assert response.headers["content-type"].startswith(MEDIA_TYPES[format])
        assert response.headers["content-disposition"].startswith("attachment;")
        assert response.headers["etag"] == '"' + hashlib.sha256(body).hexdigest() + '"'
        assert listed["content_sha256"] == hashlib.sha256(body).hexdigest()
        assert listed["size_bytes"] == len(body) and listed["attempts"] == 1
        [stored] = artifacts(live, job_id)
        assert bytes(stored["body"]) == body and stored["lease_generation"] == 1
        row = job(live, job_id)
        assert row["output_manifest"]["content_sha256"] == listed["content_sha256"]
        assert row["lease_owner"] is None and row["completed_at"] is not None
    pdf = pdf_text(bodies["PDF"])
    assert "(" + displayed + ") Tj" in pdf and digest in pdf
    sheets = xlsx_rows(bodies["XLSX"])
    values = [[value for value, _ in line] for line in sheets["Bound values"]]
    csv_rows = [line.split(",") for line in csv_text.strip().split("\r\n")]
    assert values == csv_rows
    assert [value for value, _ in sheets["Package"][-2]] == ["reconciliation_sha256", digest]
    lines = docx_text(bodies["DOCX"])
    assert "water | " + displayed in lines and "Reconciliation SHA-256 " + digest in lines

    # Rendering the same pinned package again yields the same bytes.
    worker = make_worker(live)
    tenant = live.fixture["tenant_a"]
    for format, job_id in jobs.items():
        with live.db() as c:
            pinned = c.execute("SELECT * FROM impact.report_export WHERE job_id=%s", (job_id,)).fetchone()
        twice = []
        for _ in range(2):
            with worker.transaction(tenant) as c:
                document = worker.export_model(c, tenant, pinned)
            twice.append(hashlib.sha256(render(format, document)).hexdigest())
        assert twice == [hashlib.sha256(bodies[format]).hexdigest()] * 2, format

    with live.db() as c:
        accesses = c.execute(
            "SELECT access_mode,disclosure_id,principal_id FROM impact.report_export_access WHERE job_id=ANY(%s::uuid[])",
            (list(jobs.values()),),
        ).fetchall()
        assert len(accesses) == 3 and {a["access_mode"] for a in accesses} == {"INTERNAL_DOWNLOAD"}
        items = c.execute(
            "SELECT outcome FROM impact.job_item WHERE job_id=ANY(%s::uuid[]) AND item_key='artifact'",
            (list(jobs.values()),),
        ).fetchall()
        assert [i["outcome"] for i in items] == ["CREATED"] * 3
    # Nobody else downloads through the internal route.
    path = live.path("reports", report["object_id"]) + "/exports/" + jobs["PDF"] + "/download"
    for actor in ["partner", "other_tenant"]:
        assert live.request(path, actor=actor).status_code == 404
    other = approved_report(live)
    assert (
        live.request(
            live.path("reports", other["object_id"]) + "/exports/" + jobs["PDF"] + "/download"
        ).status_code
        == 404
    )


def test_a_stale_lease_holder_cannot_complete_and_one_artifact_exists(live):
    settle(live)
    report = approved_report(live)
    job_id = request_export(live, report, "XLSX")["job_id"]
    tenant = live.fixture["tenant_a"]
    clock = Clock()
    first = make_worker(live, clock=clock, lease_seconds=30)
    second = make_worker(live, clock=clock, lease_seconds=30)
    held = next(r for r in first.claim_exports(tenant, empty_summary()) if str(r["job_id"]) == job_id)
    assert (held["lease_generation"], held["attempts"]) == (1, 1)
    assert job_id not in {str(r["job_id"]) for r in second.claim_exports(tenant, empty_summary())}
    with first.transaction(tenant) as c:
        body = render("XLSX", first.export_model(c, tenant, held))
    clock.advance(seconds=31)
    summary = empty_summary()
    taken_rows = second.claim_exports(tenant, summary)
    taken = next(r for r in taken_rows if str(r["job_id"]) == job_id)
    assert (taken["lease_generation"], taken["attempts"]) == (2, 2)
    for row in taken_rows:
        second.process_export(tenant, row, summary)
    assert summary["exports_succeeded"] >= 1
    # The first holder wakes up: its renewal and its outcome are both refused by generation.
    late = empty_summary()
    assert first.record_export(tenant, held, late, "succeeded", None, body) is False
    first.process_export(tenant, held, late)
    assert late["stale_refused"] == 2 and late["exports_succeeded"] == 0
    [stored] = artifacts(live, job_id)
    assert stored["lease_generation"] == 2
    row = job(live, job_id)
    assert (row["state"], row["lease_generation"], row["attempts"]) == ("Succeeded", 2, 2)


def test_transient_failures_back_off_then_fail_with_an_error_class_only(live, monkeypatch):
    import impact_api.worker as worker_module

    settle(live)
    report = approved_report(live)
    job_id = request_export(live, report, "DOCX")["job_id"]
    tenant = live.fixture["tenant_a"]

    def flaky(format, document):
        raise RenderError("RENDER_TIMEOUT", permanent=False)

    monkeypatch.setattr(worker_module, "render_export", flaky)
    clock = Clock()
    worker = make_worker(live, clock=clock, max_attempts=2)
    drain(worker, tenant)
    row = job(live, job_id)
    assert (row["state"], row["attempts"], row["last_error_class"]) == ("Queued", 1, "RENDER_TIMEOUT")
    assert row["lease_owner"] is None and row["completed_at"] is None
    # Backoff on the database clock: base 30 s times a factor in [0.5, 1).
    with live.db() as c:
        now = c.execute("SELECT now() AS now").fetchone()["now"]
    assert timedelta(seconds=10) < row["next_attempt_at"] - now <= timedelta(seconds=30)
    assert job_id not in {str(r["job_id"]) for r in worker.claim_exports(tenant, empty_summary())}
    clock.advance(seconds=31)
    drain(worker, tenant)
    row = job(live, job_id)
    assert (row["state"], row["attempts"], row["last_error_class"]) == ("Failed", 2, "RENDER_TIMEOUT")
    assert row["output_manifest"]["error_class"] == "RENDER_TIMEOUT"
    assert artifacts(live, job_id) == []
    listed = export_row(live, report, job_id)
    assert listed["state"] == "Failed" and listed["content_sha256"] is None
    # A failed export may be requested again; a permanent failure is Failed at once.
    monkeypatch.setattr(
        worker_module, "render_export", lambda f, d: (_ for _ in ()).throw(RenderError("RENDER_FAILED"))
    )
    again = request_export(live, report, "DOCX")["job_id"]
    drain(make_worker(live), tenant)
    row = job(live, again)
    assert (row["state"], row["attempts"], row["last_error_class"]) == ("Failed", 1, "RENDER_FAILED")
    with live.db() as c:
        item = c.execute(
            "SELECT outcome,error_code FROM impact.job_item WHERE job_id=%s AND item_key='artifact'", (again,)
        ).fetchone()
    assert (item["outcome"], item["error_code"]) == ("FAILED", "RENDER_FAILED")


def test_cancellation_before_start_is_honoured_and_nothing_is_rendered(live):
    report = approved_report(live)
    job_id = request_export(live, report, "PDF")["job_id"]
    body = cmd({"job_id": job_id}, report["revision_id"])
    path = live.path("reports", report["object_id"]) + "/actions/cancel-export"
    receipt = expect(live.request(path, method="POST", body=body), 200)
    assert receipt["business_state"] == "CancellationRequested" and receipt["job_id"] == job_id
    assert expect(live.request(path, method="POST", body=body), 200) == receipt
    # A job id of another report is not visible through this report.
    other = approved_report(live)
    expect(
        live.request(
            live.path("reports", other["object_id"]) + "/actions/cancel-export",
            method="POST",
            body=cmd({"job_id": job_id}, other["revision_id"]),
        ),
        404,
    )
    worker = make_worker(live)
    drain(worker, live.fixture["tenant_a"])
    assert job(live, job_id)["state"] == "Queued" and job(live, job_id)["attempts"] == 0
    summary = worker.run_once()
    assert summary["cancelled"] >= 1
    row = job(live, job_id)
    assert (row["state"], row["attempts"], row["lease_generation"]) == ("Cancelled", 0, 1)
    assert row["output_manifest"]["boundary"] == "BEFORE_START"
    assert artifacts(live, job_id) == []
    listed = export_row(live, report, job_id)
    assert listed["cancellation_requested"] and listed["cancellation_outcome"] == "CANCELLED_BEFORE_START"
    finished = expect(
        live.request(path, method="POST", body=cmd({"job_id": job_id}, report["revision_id"])), 409
    )
    assert finished["reason_code"] == "EXPORT_ALREADY_FINISHED"
    # Cancelled work can be requested again.
    assert request_export(live, report, "PDF")["business_state"] == "Queued"


# -- delivery --------------------------------------------------------------------------------------


def test_published_exports_reach_named_recipients_only_and_withdrawal_blocks_them(live):
    report, jobs = produced(live, ["PDF", "XLSX"])
    # A disclosure without exports remains possible.
    body = request_disclosure(live, report)[1]
    refused = expect(
        live.request(
            live.path("disclosure-requests"),
            method="POST",
            body=cmd({**body["data"], "export_formats": ["DOCX"]}),
        ),
        409,
    )
    assert refused["reason_code"] == "EXPORT_ARTIFACT_REQUIRED"
    # The server pins exact artifacts; the request cannot name them.
    forged = expect(
        live.request(
            live.path("disclosure-requests"),
            method="POST",
            body=cmd({**body["data"], "export_artifacts": []}),
        ),
        422,
    )
    assert forged["code"] == "VALIDATION_FAILED"
    workflow, _ = request_disclosure(live, report, export_formats=["PDF", "XLSX"])
    candidate = get(live, "workflows", workflow["object_id"] + "/candidate", actor="reviewer")
    pinned = {item["format"]: item for item in candidate["record"]["data"]["export_artifacts"]}
    assert {f: pinned[f]["job_id"] for f in pinned} == jobs
    approve(live, workflow)
    disclosure_id = workflow["data"]["candidate_id"]
    disclosure = get(live, "disclosures", disclosure_id, actor="reviewer")
    publish(live, report, disclosure)
    with live.db() as c:
        bound = c.execute(
            "SELECT format,job_id,content_sha256 FROM impact.report_publication_export WHERE disclosure_id=%s "
            "ORDER BY format",
            (disclosure_id,),
        ).fetchall()
    assert [(b["format"], str(b["job_id"])) for b in bound] == [("PDF", jobs["PDF"]), ("XLSX", jobs["XLSX"])]

    base = live.path("publications", disclosure_id)
    for extension, format in [("pdf", "PDF"), ("xlsx", "XLSX")]:
        response = live.request(base + "/download." + extension, actor="partner")
        assert response.status_code == 200
        assert hashlib.sha256(response.content).hexdigest() == pinned[format]["content_sha256"]
        assert response.headers["content-type"].startswith(MEDIA_TYPES[format])
        for actor in ["author", "reviewer", "other_tenant"]:
            assert live.request(base + "/download." + extension, actor=actor).status_code == 404
    assert live.request(base + "/download.docx", actor="partner").status_code == 404
    assert live.request(base + "/download.pptx", actor="partner").status_code == 404
    with live.db() as c:
        accesses = c.execute(
            "SELECT access_mode,format,membership_id FROM impact.report_export_access WHERE disclosure_id=%s",
            (disclosure_id,),
        ).fetchall()
    assert sorted(a["format"] for a in accesses) == ["PDF", "XLSX"]
    assert {(a["access_mode"], str(a["membership_id"])) for a in accesses} == {
        ("RECIPIENT_DOWNLOAD", live.fixture["member_partner"])
    }

    # A recipient without the download right gets no export bytes.
    viewing, _ = request_disclosure(live, report, allow_download=False, export_formats=["PDF"])
    approve(live, viewing)
    view_only = viewing["data"]["candidate_id"]
    publish(live, report, get(live, "disclosures", view_only, actor="reviewer"))
    assert live.request(live.path("publications", view_only) + "/view", actor="partner").status_code == 200
    assert (
        live.request(live.path("publications", view_only) + "/download.pdf", actor="partner").status_code
        == 404
    )

    action(
        live,
        "reports",
        report,
        "withdraw",
        {"reason": "Superseded by a corrected package."},
        actor="reviewer",
    )
    for extension in ["pdf", "xlsx"]:
        assert live.request(base + "/download." + extension, actor="partner").status_code == 404
    # Withdrawal never deletes or re-points the immutable artifacts.
    for job_id in jobs.values():
        assert len(artifacts(live, job_id)) == 1


def test_export_tables_are_tenant_fenced_and_artifacts_are_insert_only(live):
    report, jobs = produced(live, ["PDF"])
    tenant_a, tenant_b = live.fixture["tenant_a"], live.fixture["tenant_b"]
    with live.db() as c:
        c.execute("SET LOCAL ROLE impact_app")
        c.execute("SELECT set_config('impact.tenant_id',%s,true)", (tenant_b,))
        for table in ["report_export", "report_export_artifact", "report_publication_export"]:
            assert c.execute("SELECT count(*) AS n FROM impact." + table).fetchone()["n"] == 0
    for statement in [
        "UPDATE impact.report_export_artifact SET body=body",
        "DELETE FROM impact.report_export_artifact",
        "UPDATE impact.report_export SET attempts=0",
        "UPDATE impact.report_publication_export SET job_id=job_id",
        "SELECT * FROM impact.report_export_access",
        "SELECT * FROM impact.worker_export_tenants(now())",
    ]:
        with live.db() as c:
            c.execute("SET LOCAL ROLE impact_app")
            c.execute("SELECT set_config('impact.tenant_id',%s,true)", (tenant_a,))
            with pytest.raises(psycopg.errors.InsufficientPrivilege):
                c.execute(statement)
    with psycopg.connect(worker_dsn(), row_factory=dict_row, prepare_threshold=None) as c:
        c.execute("SET LOCAL ROLE impact_worker")
        directory = c.execute("SELECT * FROM impact.worker_export_tenants(now()+interval '1 day')").fetchall()
        assert all(set(row) == {"tenant_id", "due_exports"} for row in directory)
        assert c.execute("SELECT count(*) AS n FROM impact.report_export").fetchone()["n"] == 0
        c.execute("SELECT set_config('impact.tenant_id',%s,true)", (tenant_b,))
        assert c.execute("SELECT count(*) AS n FROM impact.report_export_artifact").fetchone()["n"] == 0
        assert (
            c.execute(
                "UPDATE impact.report_export SET attempts=attempts WHERE job_id=%s", (jobs["PDF"],)
            ).rowcount
            == 0
        )
    for statement in [
        "UPDATE impact.report_export_artifact SET body=body",
        "DELETE FROM impact.report_export_artifact",
        "INSERT INTO impact.report_publication_export SELECT * FROM impact.report_publication_export",
    ]:
        with psycopg.connect(worker_dsn(), prepare_threshold=None) as c:
            c.execute("SET LOCAL ROLE impact_worker")
            c.execute("SELECT set_config('impact.tenant_id',%s,true)", (tenant_a,))
            with pytest.raises(psycopg.errors.InsufficientPrivilege):
                c.execute(statement)


@pytest.mark.skipif(not NATIVE, reason="concurrent claims need real PostgreSQL row locks")
def test_native_two_workers_claiming_at_once_lease_each_export_once(live):
    settle(live)
    reports = [approved_report(live) for _ in range(3)]
    jobs = {request_export(live, report, "XLSX")["job_id"] for report in reports}
    tenant = live.fixture["tenant_a"]
    workers = [make_worker(live), make_worker(live)]
    barrier = threading.Barrier(2)
    claimed = [[], []]

    def run(index):
        barrier.wait()
        claimed[index] = workers[index].claim_exports(tenant, empty_summary())

    threads = [threading.Thread(target=run, args=(i,)) for i in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    mine = [{str(r["job_id"]) for r in rows} & jobs for rows in claimed]
    # SKIP LOCKED: each job is leased by exactly one of the two simultaneous claims.
    assert not (mine[0] & mine[1]) and mine[0] | mine[1] == jobs
    for job_id in jobs:
        row = job(live, job_id)
        assert (row["state"], row["lease_generation"], row["attempts"]) == ("Running", 1, 1)
    for worker, rows in zip(workers, claimed):
        summary = empty_summary()
        for row in rows:
            worker.process_export(tenant, row, summary)
        assert summary["stale_refused"] == 0
    assert all(job(live, job_id)["state"] == "Succeeded" for job_id in jobs)
    assert all(len(artifacts(live, job_id)) == 1 for job_id in jobs)
