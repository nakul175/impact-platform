"""Build the four editable specifications from their canonical Markdown sources.

Requires python-docx. Render and visually review every output before delivery.
This authoring helper has no application, database or network access.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CORE = ("01-BRD", "02-FSD", "03-HLD", "04-LLD")
LABELS = {
    "01-BRD": "Business Requirements Document",
    "02-FSD": "Functional Specification Document",
    "03-HLD": "High Level Design",
    "04-LLD": "Low Level Design",
}


def plain(value):
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    return value.replace("**", "").replace("`", "").strip()


def font(run, size=None, bold=None):
    run.font.name = "Arial"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string("000000")


def inline(paragraph, value):
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", value):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[2:-2] if bold else part[1:-1] if code else part)
        font(run, bold=bold)
        if code:
            run.font.name = "Courier New"
            run.font.size = Pt(9)


def set_cell(cell, value, header=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.line_spacing = 1.06
    inline(p, value)
    for run in p.runs:
        font(run, 9, header or run.bold)
        if header:
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "80" if side in ("left", "right") else "50")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def add_table(document, rows):
    count = len(rows[0])
    table = document.add_table(rows=0, cols=count)
    table.autofit = False
    # Bounded content weights keep identifiers narrow and descriptions readable.
    weights = []
    for col in range(count):
        lengths = sorted(len(plain(row[col])) for row in rows if len(row) > col)
        typical = lengths[len(lengths) // 2]
        weights.append(min(70, max(18, typical)))
    total = sum(weights)
    widths = [6.8 * weight / total for weight in weights]
    if count == 2 and plain(rows[0][0]) == "Code and status":
        # Error identifiers must remain intact when the document is rendered.
        widths = [6.8 * 0.32, 6.8 * 0.68]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for col, cell in enumerate(row.cells):
            cell.width = Inches(widths[col])
            set_cell(cell, values[col] if col < len(values) else "", index == 0)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "243B53" if index == 0 else "EEF3F8" if index % 2 == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def style_document(document, label):
    section = document.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.72)
    section.header_distance = section.footer_distance = Inches(0.28)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (
        ("Title", 23),
        ("Subtitle", 14),
        ("Heading 1", 15),
        ("Heading 2", 12),
        ("Heading 3", 11),
    ):
        style = document.styles[name]
        for border in list(style.element.findall(".//" + qn("w:pBdr"))):
            border.getparent().remove(border)
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(13 if name.startswith("Heading") else 4)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    font(header.add_run("NONPROFIT AI ENABLEMENT  |  " + label.upper()), 8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
    font(footer.add_run("Proposed baseline  |  Edition 1.0\t"), 8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "Nonprofit AI Enablement " + label
    document.core_properties.subject = "Proposed nonprofit AI extension specification"
    document.core_properties.author = "Impact Platform"
    document.core_properties.keywords = "nonprofit, AI, specification, proposed baseline"


def build(source, output):
    lines = source.read_text().splitlines()
    document = Document()
    style_document(document, LABELS[source.stem])
    document.add_paragraph("Nonprofit AI Enablement Platform", "Title")
    document.add_paragraph(LABELS[source.stem], "Subtitle")
    document.add_paragraph("Edition 1.0  |  5 October 2026  |  Proposed for owner review")
    document.add_paragraph("Product reference 36073f1  |  Build 0.30.0  |  Schema 35")
    document.add_paragraph(
        "This edition distinguishes current local implementation from proposed scope. It does not approve requirements, execute tests or authorise deployment."
    )
    document.add_paragraph("Contents", "Heading 1")
    for line in lines:
        if line.startswith("## "):
            document.add_paragraph(plain(line[3:]))
    document.add_page_break()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("# ") or line.startswith("Edition 1.0"):
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", cell.strip()) for cell in row):
                    rows.append(row)
                index += 1
            add_table(document, rows)
            continue
        if line.startswith("```"):
            language = line[3:]
            index += 1
            code = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if language == "mermaid":
                document.add_paragraph("Diagram source is included in the companion Markdown specification.")
                index += 1
                continue
            for code_index, value in enumerate(code):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = code_index < len(code) - 1
                run = paragraph.add_run(value)
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string("000000")
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            asset = source.parent / image_match[2]
            png = asset.with_suffix(".png") if asset.suffix == ".svg" else asset
            if png.exists():
                document.add_picture(str(png), width=Inches(6.8))
                p = document.add_paragraph(image_match[1])
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    font(run, 9)
            else:
                document.add_paragraph(image_match[1] + " is supplied in the companion package.")
            index += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            # Keep meaningful requirement identifiers; strip decorative punctuation only.
            text = plain(heading[2]).replace(" -> ", " to ").replace(" / ", " and ")
            document.add_paragraph(text, "Heading " + str(len(heading[1]) - 1))
        elif line.startswith(("- ", "* ")):
            p = document.add_paragraph(style="List Bullet")
            inline(p, line[2:])
        else:
            p = document.add_paragraph()
            inline(p, line)
            if re.match(r"\*\*(BR|FR)-NPA-", line):
                p.paragraph_format.keep_with_next = True
        index += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path("docs/nonprofit-ai/v1.0"))
    parser.add_argument("--output", type=Path)
    parser.add_argument("--only", nargs="+", choices=CORE)
    args = parser.parse_args()
    for name in args.only or CORE:
        source = args.source / f"{name}.md"
        output = (args.output or args.source / "editable") / f"{name}.docx"
        build(source, output)
        print(output)


if __name__ == "__main__":
    main()
